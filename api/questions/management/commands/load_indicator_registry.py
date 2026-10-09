import json
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from api.interviews.management.commands.import_governance_question_bank import _split_points
from api.questions.models import IndicatorDefinition

BANK_VERSION = "GOV1.2-FINAL"
PRIORITY_NINE = (
    "CCG-TSK-002", "CC-TEK-005", "ECG-TSK-001", "FDA-TSK-004", "HK-TSK-001",
    "IC-TEK-004", "CCG-TSK-003", "IC-TSK-003", "SNC-GEN-009",
)
DEFAULT_FIXTURE = (
    Path(settings.BASE_DIR) / "api" / "interviews" / "fixtures" / "question_bank" / "governance_v1_2_corrected.json"
)


def _comparable(text):
    return " ".join(str(text or "").strip().rstrip(".").lower().split())


def _read_arabic_drafts(path):
    """{indicator_id: (english, arabic)} from a Priority 9 package sheet with
    'Indicator ID', 'English ...' and 'Arabic ...' columns."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(rows, [])]
        if "Indicator ID" not in header:
            continue
        id_col = header.index("Indicator ID")
        en_col = next((i for i, h in enumerate(header) if h.startswith("English")), None)
        ar_col = next((i for i, h in enumerate(header) if h.startswith("Arabic")), None)
        if en_col is None or ar_col is None:
            continue
        drafts = {}
        for row in rows:
            if row and row[id_col]:
                drafts[str(row[id_col]).strip()] = (row[en_col], row[ar_col])
        return drafts, sheet.title
    raise CommandError(f"No sheet with Indicator ID / English / Arabic columns in {path}")


class Command(BaseCommand):
    help = (
        "Create the canonical indicator-ID registry for the approved bank (all 441 questions, or "
        "--priority-nine) with locked English. Optionally attach DRAFT Arabic text from a package "
        "workbook - only where ALLOW_DRAFT_INDICATOR_TRANSLATIONS is enabled (Staging, policy "
        "D-01). Never changes existing English text; re-running is safe."
    )

    def add_arguments(self, parser):
        parser.add_argument("--arabic-draft", help="Workbook with draft Arabic per Indicator ID (Staging only).")
        parser.add_argument("--fixture", default=str(DEFAULT_FIXTURE))
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument("--priority-nine", action="store_true",
                            help="Limit the registry to the nine Priority 9 questions (default: whole approved bank).")

    def handle(self, *args, **options):
        if options["arabic_draft"] and not settings.ALLOW_DRAFT_INDICATOR_TRANSLATIONS:
            raise CommandError(
                "Draft Arabic indicator text may only be loaded where ALLOW_DRAFT_INDICATOR_TRANSLATIONS "
                "is enabled (Staging). Policy D-01: no draft Arabic mapping in Production."
            )

        bank = {q["question_code"]: q for q in json.loads(Path(options["fixture"]).read_text(encoding="utf-8"))["questions"]}
        missing = [code for code in PRIORITY_NINE if code not in bank]
        if missing:
            raise CommandError(f"Priority 9 codes missing from the bank: {missing}")
        codes = PRIORITY_NINE if options["priority_nine"] else list(bank)

        planned = []
        for code in codes:
            q = bank[code]
            for kind, source, prefix in (
                (IndicatorDefinition.TYPE_MUST_INCLUDE, q["must_include_en"], "MI"),
                (IndicatorDefinition.TYPE_NEGATIVE, q["negative_indicators"], "NI"),
            ):
                for ordinal, text in enumerate(_split_points(source), start=1):
                    planned.append((f"{code}-{prefix}-{ordinal:02d}", code, kind, ordinal, text))

        drafts, sheet_name = ({}, "")
        if options["arabic_draft"]:
            drafts, sheet_name = _read_arabic_drafts(options["arabic_draft"])
            unknown = sorted(set(drafts) - {p[0] for p in planned})
            if unknown:
                raise CommandError(f"Draft workbook has IDs not in the registry: {unknown[:5]}")
            mismatched = [
                indicator_id for indicator_id, _, _, _, text in planned
                if indicator_id in drafts and _comparable(drafts[indicator_id][0]) != _comparable(text)
            ]
            if mismatched:
                raise CommandError(f"Draft workbook English differs from the locked bank for: {mismatched}")

        conflicts = [
            indicator_id for indicator_id, _, _, _, text in planned
            if IndicatorDefinition.objects.filter(indicator_id=indicator_id).exclude(text_en=text).exists()
        ]
        if conflicts:
            raise CommandError(f"Existing indicators would change English text (IDs are immutable): {conflicts}")

        self.stdout.write(
            f"{len(planned)} indicators planned "
            f"({sum(1 for p in planned if p[2] == IndicatorDefinition.TYPE_MUST_INCLUDE)} Must Include, "
            f"{sum(1 for p in planned if p[2] == IndicatorDefinition.TYPE_NEGATIVE)} Negative); "
            f"draft Arabic for {len(drafts)}."
        )
        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("Dry run: nothing written."))
            return

        with transaction.atomic():
            for indicator_id, code, kind, ordinal, text in planned:
                record, _ = IndicatorDefinition.objects.get_or_create(
                    indicator_id=indicator_id,
                    defaults={"question_code": code, "bank_version": BANK_VERSION, "indicator_type": kind,
                              "ordinal": ordinal, "text_en": text},
                )
                if indicator_id in drafts and record.text_ar_status != IndicatorDefinition.STATUS_APPROVED:
                    record.text_ar = str(drafts[indicator_id][1] or "").strip()
                    record.text_ar_status = IndicatorDefinition.STATUS_DRAFT
                    record.text_ar_source = f"{Path(options['arabic_draft']).name} / {sheet_name}"[:120]
                    record.save(update_fields=["text_ar", "text_ar_status", "text_ar_source", "updated_at"])
        self.stdout.write(self.style.SUCCESS(f"Registry ready: {len(planned)} indicators."))
