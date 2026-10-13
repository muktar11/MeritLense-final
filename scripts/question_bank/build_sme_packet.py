"""Build the SME review workbook for GOV2.0-RC1: every AI-drafted or
corrected item awaiting specialist / bilingual sign-off, priority sectors
first, with decision columns for the reviewer.

Usage: python scripts/question_bank/build_sme_packet.py <out.xlsx>
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[2]
PACKAGE = REPO / "api/interviews/fixtures/question_bank/governance_v2_0.json"
CORRECTIONS = REPO / "scripts/question_bank/corrections"
PRIORITY_SECTORS = {
    "Nursing Assistant": "Clinical / care", "Home Care Assistant": "Clinical / care",
    "Elderly Medical Support": "Clinical / care", "Basic Patient Support": "Clinical / care",
    "Special Needs Caregiver": "Clinical / care", "Elderly Caregiver": "Clinical / care",
    "Child Caregiver": "Child safety", "Restaurant Staff": "Food hygiene",
    "Industrial Cleaner": "Industrial hazards", "Skilled Trades": "Trade hazards",
    "General Labor": "Trade hazards", "Farm Worker": "Pesticides / food hygiene",
    "Livestock Support": "Animal health",
}
DECISION = ["SME decision (Approve / Amend / Reject)", "SME comment", "Reviewer", "Date"]
HEADER_FILL = PatternFill("solid", fgColor="1D5F78")


def _sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers + DECISION)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append(row + ["", "", "", ""])
    for i, width in enumerate(widths + [22, 30, 14, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    return ws


def main(out):
    package = json.loads(PACKAGE.read_text(encoding="utf-8"))
    qs = package["questions"]

    def priority(q):
        return (0 if q["role"] in PRIORITY_SECTORS else 1, q["criticality"] != "Critical", q["role"], q["question_code"])

    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    for line in [
        "MeritLense GOV2.0-RC1 - SME review packet",
        "Every item here was drafted or corrected by AI editorial review and is NOT approved.",
        "Priority sectors first: clinical/care, child safety, food hygiene, industrial/trade hazards, pesticides, animal health.",
        "Severity (D-04): S3 critical safety, S2 major, S1 noncritical. Mandatory (D-02): points required for full marks; other points are scored, not removed.",
        "D-06 questions (CCG-TSK-003, IC-TSK-003, SNC-GEN-009) are kept at the approved v1.2 content; the handoff rewrite is listed for review.",
        "Record a decision per row. Nothing changes in Production until management approves the closure items.",
    ]:
        readme.append([line])
    readme.column_dimensions["A"].width = 140

    neg = sorted((q for q in qs if "NEGATIVE_INDICATORS" in q.get("corrections_applied", [])), key=priority)
    _sheet(wb, "Negative indicators", ["Question", "Role", "Priority sector", "Competency", "Criticality",
                                       "Question (EN)", "Must Include (EN)", "Proposed negative indicators", "Proposed severities"],
           [[q["question_code"], q["role"], PRIORITY_SECTORS.get(q["role"], ""), q["competency_en"], q["criticality"],
             q["question_en"], q["must_include_en"], q["negative_indicators"], ", ".join(q.get("negative_severity", []))]
            for q in neg], [14, 18, 16, 18, 11, 45, 55, 55, 14])

    mand = sorted((q for q in qs if q.get("mandatory_indicators")), key=priority)
    _sheet(wb, "Mandatory points (D-02)", ["Question", "Role", "Priority sector", "Criticality", "Question (EN)",
                                           "All Must Include points", "Proposed mandatory points"],
           [[q["question_code"], q["role"], PRIORITY_SECTORS.get(q["role"], ""), q["criticality"], q["question_en"],
             q["must_include_en"], "; ".join(q["mandatory_indicators"])] for q in mand], [14, 18, 16, 11, 45, 60, 45])

    alignment = json.loads((CORRECTIONS / "03_alignment.json").read_text(encoding="utf-8"))["corrections"]
    by_code = {q["question_code"]: q for q in qs}
    _sheet(wb, "Arabic corrections", ["Question", "Role", "Kind", "Finding", "English", "Corrected Arabic"],
           [[c, by_code[c]["role"], v["kind"], v["note"],
             by_code[c]["must_include_en"] if "must_include_ar" in v["fields"] else by_code[c]["question_en"],
             v["fields"].get("must_include_ar") or v["fields"].get("question_ar")]
            for c, v in sorted(alignment.items())], [14, 18, 9, 40, 55, 60])

    dup = json.loads((CORRECTIONS / "04_duplicate_replacements.json").read_text(encoding="utf-8"))["corrections"]
    _sheet(wb, "Replacement questions", ["Question", "Replaces", "Question (EN)", "Question (AR)", "Must Include (EN)",
                                         "Must Include (AR)", "Negative indicators", "Severities"],
           [[c, v["replaces"], v["fields"]["question_en"], v["fields"]["question_ar"], v["fields"]["must_include_en"],
             v["fields"]["must_include_ar"], v["fields"]["negative_indicators"], ", ".join(v["fields"]["negative_severity"])]
            for c, v in dup.items()], [12, 40, 45, 45, 50, 50, 45, 12])

    prompts = json.loads((CORRECTIONS / "02_missing_prompts.json").read_text(encoding="utf-8"))["corrections"]
    _sheet(wb, "Prompt fixes", ["Question", "Role", "Question (EN) now", "Question (AR) now"],
           [[c, by_code[c]["role"], by_code[c]["question_en"], by_code[c]["question_ar"]] for c in sorted(prompts)],
           [14, 18, 60, 60])

    d06 = json.loads((CORRECTIONS / "01_d06_locked.json").read_text(encoding="utf-8"))["corrections"]
    _sheet(wb, "D-06 held proposals", ["Question", "Field", "Approved v1.2 (kept)", "Handoff proposal (held)"],
           [[c, f, by_code[c][f], v["proposed_for_sme"][f]] for c, v in d06.items() for f in v["proposed_for_sme"]],
           [14, 18, 60, 60])

    matrix = json.loads((CORRECTIONS / "07_allocation_matrix.json").read_text(encoding="utf-8"))["roles"]
    _sheet(wb, "Selection matrix", ["Role", "Full allocation", "Screening allocation", "Rationale"],
           [[r, json.dumps(m["FULL"]), json.dumps(m["SCREENING"]), m["rationale"]] for r, m in matrix.items()],
           [22, 45, 40, 70])

    wb.save(out)
    print(f"Wrote {out}: {len(neg)} negative sets, {len(mand)} mandatory sets, {len(alignment)} Arabic corrections, "
          f"{len(dup)} replacements, {len(prompts)} prompt fixes, {len(d06)} D-06 holds, {len(matrix)} roles")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "MeritLense_Question_Bank_v2_SME_Review_Packet.xlsx")
